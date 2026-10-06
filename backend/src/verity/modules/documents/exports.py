"""The acknowledgement completion export: one campaign, as CSV or Excel.

It is audit evidence for the question "how do you know staff read the policy", so it
carries the document, the version and the campaign beside every signature and says
who exported it and when. The CSV is a clean table (one header row, one row per
person, the document and campaign columns repeated) because anything else breaks a
pivot or an import; the Excel file holds the same table plus a second sheet for the
facts about the export itself.

Every text cell passes through :func:`safe_cell`. A cell that starts with a character a
spreadsheet reads as a formula would run when someone opens the file, and these cells
hold names, addresses and free-text comments typed by other people.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import Final, Literal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ExportFormat = Literal["csv", "xlsx"]

CONTENT_TYPES: Final[dict[str, str]] = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

HEADERS: Final = (
    "Document code",
    "Document title",
    "Version",
    "Campaign",
    "Due date",
    "Name",
    "Email",
    "Kind",
    "Status",
    "Acknowledged at",
    "Comment",
)

_FORMULA_STARTS: Final = ("=", "+", "-", "@", "\t", "\r")
_HEADER_FILL: Final = "1F2A44"


@dataclass(frozen=True, slots=True)
class ExportRecipient:
    name: str
    email: str
    kind: str
    signed: bool
    acknowledged_at: datetime | None
    comment: str | None


@dataclass(frozen=True, slots=True)
class AckExport:
    """Everything the file says, gathered by the service so rendering is pure."""

    document_code: str
    document_title: str
    version: str | None
    campaign_title: str
    campaign_status: str
    due_at: datetime | None
    exported_by: str
    exported_at: datetime
    recipients: list[ExportRecipient] = field(default_factory=list)


def safe_cell(value: object) -> str:
    """``value`` as cell text, with a leading apostrophe if a spreadsheet would
    otherwise take it for a formula."""
    text = "" if value is None else str(value)
    return f"'{text}" if text.startswith(_FORMULA_STARTS) else text


def _stamp(moment: datetime | None) -> str:
    return moment.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ") if moment else ""


def table(export: AckExport) -> list[list[str]]:
    """The header row and one row per recipient, every cell already safe."""
    due = export.due_at.astimezone(UTC).date().isoformat() if export.due_at else ""
    rows = [list(HEADERS)]
    for person in export.recipients:
        rows.append(
            [
                safe_cell(cell)
                for cell in (
                    export.document_code,
                    export.document_title,
                    export.version,
                    export.campaign_title,
                    due,
                    person.name,
                    person.email,
                    person.kind,
                    "Signed" if person.signed else "Pending",
                    _stamp(person.acknowledged_at),
                    person.comment,
                )
            ]
        )
    return rows


def _facts(export: AckExport) -> list[tuple[str, str]]:
    signed = sum(1 for p in export.recipients if p.signed)
    return [
        ("Exported by", export.exported_by),
        ("Exported at", _stamp(export.exported_at)),
        ("Document", f"{export.document_code} {export.document_title}"),
        ("Version", export.version or "Not set"),
        ("Campaign", export.campaign_title),
        ("Campaign status", export.campaign_status.capitalize()),
        ("People asked", str(len(export.recipients))),
        ("Signed", str(signed)),
        ("Pending", str(len(export.recipients) - signed)),
    ]


def render_csv(export: AckExport) -> bytes:
    buffer = io.StringIO()
    csv.writer(buffer).writerows(table(export))
    # A byte order mark so Excel reads names with accents as UTF-8.
    return buffer.getvalue().encode("utf-8-sig")


def render_xlsx(export: AckExport) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None  # noqa: S101 - a new workbook always has its first sheet
    sheet.title = "Acknowledgements"
    for row in table(export):
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)
    for index, width in enumerate((14, 36, 10, 36, 12, 24, 32, 10, 10, 22, 48), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"

    about = workbook.create_sheet("Export")
    for label, value in _facts(export):
        about.append([label, safe_cell(value)])
    for cell in about["A"]:
        cell.font = Font(bold=True)
    about.column_dimensions["A"].width = 18
    about.column_dimensions["B"].width = 60

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def filename(export: AckExport, file_format: str) -> str:
    """``POL-01-acknowledgements-2026-10-06.csv``: ASCII only, safe in a header."""
    code = re.sub(r"[^A-Za-z0-9]+", "-", export.document_code).strip("-") or "document"
    stamp = export.exported_at.astimezone(UTC).date().isoformat()
    return f"{code}-acknowledgements-{stamp}.{file_format}"
