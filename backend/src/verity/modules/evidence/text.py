"""The words in an evidence file, for the AI to read. Best effort, bounded, never raises.

Plain text, JSON, CSV, Word and Excel are read; anything else (a scan, a PDF, an image)
returns nothing and the model is told the content was not readable. The text is DATA in a
prompt, never instructions (rule 8).
"""

from __future__ import annotations

import io
import re
import zipfile

MAX_CHARS = 12_000
_MAX_BYTES = 5_000_000
_TAGS = re.compile(r"<[^>]+>")


def extract_text(data: bytes, content_type: str, filename: str) -> str:  # noqa: PLR0911
    if len(data) > _MAX_BYTES:
        return ""
    name = filename.lower()
    try:
        if content_type.startswith("text/") or content_type == "application/json":
            return data.decode("utf-8", errors="replace")[:MAX_CHARS]
        if name.endswith(".docx") or "wordprocessingml" in content_type:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
            return _TAGS.sub(" ", xml.replace("</w:p>", "\n"))[:MAX_CHARS]
        if name.endswith(".xlsx") or "spreadsheetml" in content_type:
            from openpyxl import load_workbook  # noqa: PLC0415

            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            lines: list[str] = []
            for sheet in book.worksheets[:5]:
                for row in sheet.iter_rows(max_row=200, values_only=True):
                    lines.append(" | ".join("" if c is None else str(c) for c in row))
                    if sum(len(x) for x in lines) > MAX_CHARS:
                        return "\n".join(lines)[:MAX_CHARS]
            return "\n".join(lines)[:MAX_CHARS]
    except Exception:
        return ""
    return ""
