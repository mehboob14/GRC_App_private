"""The audit export's cell rules, CSV stream, workbook and window — no database.

The route, the batches and the audit row it writes are proved against Postgres in
``tests/integration/test_audit_export.py``; what is here is everything that can be
decided from the values alone.
"""

from __future__ import annotations

import codecs
import csv
import io
import uuid
from collections.abc import AsyncIterator
from datetime import UTC, date, datetime
from typing import Any

import pytest
from openpyxl import load_workbook

from verity.core.errors import InvalidInput
from verity.modules.audit.export import (
    HEADERS,
    build_xlsx,
    csv_stream,
    events,
    neutralise_cell,
    row_cells,
)
from verity.modules.audit.models import AuditLog
from verity.modules.audit.schemas import AuditLogEntry
from verity.modules.audit.service import (
    AuditLabels,
    AuditService,
    ExportBatch,
    ExportRequest,
    ExportRun,
    Membership,
    _export_window,
)
from verity.shared.ids import uuid7

MOMENT = datetime(2026, 3, 15, 10, 30, 5, 123456, tzinfo=UTC)


def _entry(**overrides: Any) -> AuditLogEntry:  # noqa: ANN401 — a field of the model
    values: dict[str, Any] = {
        "id": uuid7(),
        "tenant_id": uuid7(),
        "actor_type": "membership",
        "actor_id": uuid7(),
        "action": "update",
        "object_type": "role",
        "object_id": uuid7(),
        "before": {"name": "Old"},
        "after": {"name": "New"},
        "occurred_at": MOMENT,
    }
    values.update(overrides)
    return AuditLogEntry(**values)


async def _batches(*batches: list[AuditLogEntry]) -> AsyncIterator[list[AuditLogEntry]]:
    for batch in batches:
        yield batch


def _run(**overrides: Any) -> ExportRun:  # noqa: ANN401 — a field of the run
    values: dict[str, Any] = {
        "export_id": uuid7(),
        "tenant_id": uuid7(),
        "request": ExportRequest("xlsx", date(2026, 3, 1), None, True, "role"),
        "cut_off": datetime(2026, 10, 6, 12, 0, tzinfo=UTC),
        "actor_label": "Ada Lovelace",
    }
    values.update(overrides)
    return ExportRun(**values)


# ---------------------------------------------------------------------------
# The cell rule
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "value",
    ["=1+1", "+1", "-1", "@SUM(A1)", "\tcmd", "\rcmd", '=HYPERLINK("http://x","y")'],
)
def test_a_cell_a_spreadsheet_would_run_gets_a_leading_apostrophe(value: str) -> None:
    assert neutralise_cell(value) == f"'{value}"


@pytest.mark.parametrize(
    "value",
    ["", "plain", "a=b", "x-y", " =1+1", "'=1+1", "1+1", '{"a":1}', "José"],
)
def test_ordinary_text_is_left_alone(value: str) -> None:
    assert neutralise_cell(value) == value


def test_every_column_of_a_row_passes_through_the_rule() -> None:
    entry = _entry(actor_label="=cmd", object_label="@sum", object_type="-kind")

    cells = row_cells(entry)

    assert len(cells) == len(HEADERS)
    assert cells[HEADERS.index("actor")] == "'=cmd"
    assert cells[HEADERS.index("object")] == "'@sum"
    assert cells[HEADERS.index("object_type")] == "'-kind"


def test_a_row_reads_as_the_api_does() -> None:
    entry = _entry(before={"b": 1, "a": {"z": True, "y": None}}, after=None)

    cells = dict(zip(HEADERS, row_cells(entry), strict=True))

    assert cells["occurred_at"] == "2026-03-15T10:30:05.123456Z"
    assert cells["actor"] == "Membership", "the fallback label, as in the list"
    assert cells["actor_id"] == str(entry.actor_id)
    assert cells["object"] == "Role"
    assert cells["object_id"] == str(entry.object_id)
    assert cells["before"] == '{"a":{"y":null,"z":true},"b":1}', "compact, keys sorted"
    assert cells["after"] == "", "no snapshot is an empty cell"


def test_a_system_row_has_no_actor_id_cell() -> None:
    cells = dict(zip(HEADERS, row_cells(_entry(actor_type="system", actor_id=None)), strict=True))

    assert cells["actor"] == "System"
    assert cells["actor_id"] == ""


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------


async def test_csv_streams_the_header_and_then_one_chunk_per_batch() -> None:
    batches = [[_entry(), _entry()], [_entry()]]

    chunks = [chunk async for chunk in csv_stream(_batches(*batches))]

    assert len(chunks) == 3
    assert chunks[0] == codecs.BOM_UTF8 + ",".join(HEADERS).encode() + b"\r\n"
    rows = list(csv.reader(io.StringIO(b"".join(chunks).decode("utf-8-sig"))))
    assert rows[0] == list(HEADERS)
    assert len(rows) == 1 + 3


async def test_csv_quotes_commas_quotes_and_line_breaks_so_they_read_back_whole() -> None:
    awkward = 'Say "hi", then\nleave'
    entry = _entry(object_label=awkward, before={"note": 'a,b "c"'})

    text = b"".join([chunk async for chunk in csv_stream(_batches([entry]))]).decode("utf-8-sig")

    (_, row) = list(csv.reader(io.StringIO(text)))
    cells = dict(zip(HEADERS, row, strict=True))
    assert cells["object"] == awkward
    assert cells["before"] == '{"note":"a,b \\"c\\""}'


async def test_csv_of_nothing_is_just_the_header() -> None:
    chunks = [chunk async for chunk in csv_stream(_batches())]

    assert chunks == [codecs.BOM_UTF8 + ",".join(HEADERS).encode() + b"\r\n"]


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def _open(data: bytes) -> Any:  # noqa: ANN401 — openpyxl is untyped
    return load_workbook(io.BytesIO(data))


async def test_the_workbook_has_the_events_then_the_exports_own_details() -> None:
    batches = [[_entry(), _entry()], [_entry()]]
    run = _run()

    book = _open(await build_xlsx(_batches(*batches), run))

    assert book.sheetnames == ["Audit log", "Export"]
    events = book["Audit log"]
    assert [cell.value for cell in events[1]] == list(HEADERS)
    assert events.max_row == 1 + 3
    assert events.freeze_panes == "A2"
    assert events.auto_filter.ref == "A1:J4", "the filter covers the rows, counted at the end"
    assert events["A1"].font.bold
    assert events["A1"].fill.fgColor.rgb.endswith("1F2A44")
    assert events.column_dimensions["I"].width == 50
    details = {row[0]: row[1] for row in book["Export"].iter_rows(values_only=True)}
    assert details == {
        "Export id": str(run.export_id),
        "Exported by": "Ada Lovelace",
        "Exported at (UTC)": "2026-10-06T12:00:00Z",
        "Cut-off (UTC)": "2026-10-06T12:00:00Z",
        "From (UTC day)": "2026-03-01",
        "To (UTC day)": "Not set",
        "Include system activity": "Yes",
        "Object type": "role",
        "Rows": 3,
    }


async def test_a_workbook_of_nothing_still_opens() -> None:
    book = _open(await build_xlsx(_batches(), _run(request=ExportRequest("xlsx"))))

    assert book["Audit log"].max_row == 1
    assert book["Audit log"].auto_filter.ref == "A1:J1"
    details = {row[0]: row[1] for row in book["Export"].iter_rows(values_only=True)}
    assert (details["Rows"], details["Object type"], details["From (UTC day)"]) == (
        0,
        "All",
        "Not set",
    )


async def test_a_formula_is_stored_as_text_everywhere_in_the_workbook() -> None:
    entry = _entry(actor_label="=1+1", object_label="@x")
    run = _run(actor_label="=Eve", request=ExportRequest("xlsx", object_type="=cmd"))

    book = _open(await build_xlsx(_batches([entry]), run))

    row = book["Audit log"][2]
    assert (row[1].value, row[1].data_type) == ("'=1+1", "s")
    assert (row[6].value, row[6].data_type) == ("'@x", "s")
    details = {row[0].value: row[1] for row in book["Export"].iter_rows()}
    assert (details["Exported by"].value, details["Exported by"].data_type) == ("'=Eve", "s")
    assert details["Object type"].value == "'=cmd"


async def test_a_stray_control_character_cannot_fail_the_whole_file() -> None:
    """openpyxl raises on these, and a write-only sheet is left half written by it."""
    entry = _entry(object_label="bad\x01name\x0b", actor_label="ok")

    book = _open(await build_xlsx(_batches([entry, _entry()]), _run()))

    assert book["Audit log"][2][6].value == "badname"
    assert book["Audit log"].max_row == 3


async def test_a_cell_over_excels_limit_is_cut_and_says_so_while_csv_keeps_it_whole() -> None:
    entry = _entry(before={"note": "x" * 40_000})
    full = row_cells(entry)[HEADERS.index("before")]

    book = _open(await build_xlsx(_batches([entry]), _run()))

    cut = book["Audit log"][2][8].value
    assert len(cut) == 32_767
    assert cut.endswith("[cut: export CSV for the full text]")
    assert cut.startswith('{"note":"xxx')
    assert len(full) > 40_000
    text = b"".join([chunk async for chunk in csv_stream(_batches([entry]))]).decode("utf-8-sig")
    (_, csv_row) = list(csv.reader(io.StringIO(text)))
    assert csv_row[HEADERS.index("before")] == full, "nothing is cut in a CSV"


# ---------------------------------------------------------------------------
# The window and the request
# ---------------------------------------------------------------------------


def test_the_window_holds_both_days_whole_and_leaves_the_cut_off_out() -> None:
    tenant, cut_off = uuid7(), datetime(2026, 10, 6, tzinfo=UTC)

    window = _export_window(
        tenant, ExportRequest("csv", date(2026, 3, 1), date(2026, 3, 2)), cut_off
    )

    assert window.tenant_id == tenant
    assert window.occurred_from == datetime(2026, 3, 1, 0, 0, tzinfo=UTC)
    assert window.occurred_to == datetime(2026, 3, 2, 23, 59, 59, 999_999, tzinfo=UTC)
    assert window.occurred_before == cut_off
    assert window.exclude_object_types, "auth telemetry is out unless asked for"


def test_an_open_window_has_no_date_bounds_and_system_activity_can_be_included() -> None:
    cut_off = datetime(2026, 10, 6, tzinfo=UTC)

    window = _export_window(
        uuid7(), ExportRequest("csv", include_system=True, object_type="x"), cut_off
    )

    assert (window.occurred_from, window.occurred_to) == (None, None)
    assert window.exclude_object_types is None
    assert window.object_type == "x"


def test_the_last_day_there_is_does_not_overflow() -> None:
    window = _export_window(
        uuid7(), ExportRequest("csv", date.min, date.max), datetime(2026, 10, 6, tzinfo=UTC)
    )

    assert window.occurred_to == datetime(9999, 12, 31, 23, 59, 59, 999_999, tzinfo=UTC)


async def test_a_range_that_runs_backwards_is_refused_before_anything_is_read() -> None:
    request = ExportRequest("csv", date(2026, 3, 2), date(2026, 3, 1))

    with pytest.raises(InvalidInput, match="From date is after the To date"):
        await AuditService().begin_export(
            tenant_id=uuid7(), actor=Membership(uuid7()), request=request
        )


# ---------------------------------------------------------------------------
# Labels, shared by the list and the export
# ---------------------------------------------------------------------------


def _row(actor: uuid.UUID, role: uuid.UUID) -> AuditLog:
    return AuditLog(
        id=uuid7(),
        tenant_id=uuid7(),
        actor_type="membership",
        actor_id=actor,
        action="update",
        object_type="role",
        object_id=role,
        before=None,
        after=None,
        occurred_at=MOMENT,
    )


def test_a_resolved_name_replaces_the_fallback_label_and_nothing_else_does() -> None:
    actor, role = uuid7(), uuid7()
    row = _row(actor, role)

    named = AuditLogEntry.labelled(
        row, AuditLabels({actor: "Ada Lovelace"}, {}, {role: "Auditor"}, {}, {}, {}, {})
    )
    unnamed = AuditLogEntry.labelled(row, AuditLabels({}, {}, {}, {}, {}, {}, {}))

    assert (named.actor_label, named.object_label) == ("Ada Lovelace", "Auditor")
    assert (unnamed.actor_label, unnamed.object_label) == ("Membership", "Role")
    assert named.occurred_at == MOMENT


async def test_events_turns_each_batch_into_the_events_the_list_shows() -> None:
    actor, role = uuid7(), uuid7()
    names = AuditLabels({actor: "Ada Lovelace"}, {}, {role: "Auditor"}, {}, {}, {}, {})

    async def batches() -> AsyncIterator[ExportBatch]:
        yield [_row(actor, role)], names
        yield [_row(actor, uuid7()), _row(uuid7(), role)], names

    shown = [[(e.actor_label, e.object_label) for e in batch] async for batch in events(batches())]

    assert shown == [
        [("Ada Lovelace", "Auditor")],
        [("Ada Lovelace", "Role"), ("Membership", "Auditor")],
    ]
