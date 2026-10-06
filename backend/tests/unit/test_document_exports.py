"""The acknowledgement completion export: its table, its formula guard, its two formats.

The guard matters more than it looks. Names, addresses and comments are typed by other
people, and a cell that begins with ``=`` runs when someone opens the file.
"""

from __future__ import annotations

import csv
import io
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from verity.modules.documents import exports

EXPORTED_AT = datetime(2026, 10, 6, 9, 30, 15, tzinfo=UTC)
SIGNED_AT = datetime(2026, 10, 2, 14, 5, 0, tzinfo=UTC)


def _export(**overrides: object) -> exports.AckExport:
    fields: dict[str, object] = {
        "document_code": "POL-01",
        "document_title": "Acceptable Use Policy",
        "version": "1.2",
        "campaign_title": "Annual acknowledgement",
        "campaign_status": "active",
        "due_at": datetime(2026, 10, 20, tzinfo=UTC),
        "exported_by": "Ada Admin (ada@example.test)",
        "exported_at": EXPORTED_AT,
        "recipients": [
            exports.ExportRecipient(
                name="Alice Reader",
                email="alice@example.test",
                kind="reviewer",
                signed=True,
                acknowledged_at=SIGNED_AT,
                comment="Read it twice.",
            ),
            exports.ExportRecipient(
                name="Bob Pending",
                email="bob@example.test",
                kind="approver",
                signed=False,
                acknowledged_at=None,
                comment=None,
            ),
        ],
    }
    fields.update(overrides)
    return exports.AckExport(**fields)  # type: ignore[arg-type]


@pytest.mark.parametrize("lead", ["=", "+", "-", "@", "\t", "\r"])
def test_a_cell_a_spreadsheet_would_run_gets_a_leading_apostrophe(lead: str) -> None:
    assert exports.safe_cell(f"{lead}HYPERLINK(1)") == f"'{lead}HYPERLINK(1)"


@pytest.mark.parametrize("text", ["plain", "a=b", "x+y", "Name - Role", " =padded", ""])
def test_ordinary_text_is_left_alone(text: str) -> None:
    assert exports.safe_cell(text) == text


def test_nothing_and_numbers_become_text() -> None:
    assert exports.safe_cell(None) == ""
    assert exports.safe_cell(42) == "42"


def test_the_table_is_one_header_row_and_one_row_per_recipient() -> None:
    rows = exports.table(_export())
    assert rows[0] == list(exports.HEADERS)
    assert rows[1:] == [
        [
            "POL-01",
            "Acceptable Use Policy",
            "1.2",
            "Annual acknowledgement",
            "2026-10-20",
            "Alice Reader",
            "alice@example.test",
            "reviewer",
            "Signed",
            "2026-10-02T14:05:00Z",
            "Read it twice.",
        ],
        [
            "POL-01",
            "Acceptable Use Policy",
            "1.2",
            "Annual acknowledgement",
            "2026-10-20",
            "Bob Pending",
            "bob@example.test",
            "approver",
            "Pending",
            "",
            "",
        ],
    ]


def test_a_campaign_with_no_due_date_or_version_leaves_those_cells_blank() -> None:
    row = exports.table(_export(due_at=None, version=None))[1]
    assert row[2] == ""
    assert row[4] == ""


def test_every_text_cell_is_guarded_not_only_the_comment() -> None:
    export = _export(
        document_title="=SUM(A1)",
        campaign_title="+cmd",
        recipients=[
            exports.ExportRecipient(
                name="-Mallory",
                email="@evil.test",
                kind="reviewer",
                signed=True,
                acknowledged_at=SIGNED_AT,
                comment='=HYPERLINK("http://evil.test","click")',
            )
        ],
    )
    row = exports.table(export)[1]
    assert row[1] == "'=SUM(A1)"
    assert row[3] == "'+cmd"
    assert row[5] == "'-Mallory"
    assert row[6] == "'@evil.test"
    assert row[10] == '\'=HYPERLINK("http://evil.test","click")'


def test_the_csv_is_a_clean_table_that_excel_reads_as_utf8() -> None:
    data = exports.render_csv(_export())
    assert data.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows == exports.table(_export())
    assert {len(row) for row in rows} == {len(exports.HEADERS)}


def test_the_csv_survives_commas_quotes_and_newlines_in_a_comment() -> None:
    export = _export(
        recipients=[
            exports.ExportRecipient(
                name="Zoë Ünal",
                email="zoe@example.test",
                kind="reviewer",
                signed=True,
                acknowledged_at=SIGNED_AT,
                comment='Fine, "mostly"\nsecond line',
            )
        ]
    )
    rows = list(csv.reader(io.StringIO(exports.render_csv(export).decode("utf-8-sig"))))
    assert rows[1][5] == "Zoë Ünal"
    assert rows[1][10] == 'Fine, "mostly"\nsecond line'


def test_the_excel_file_has_the_table_then_a_sheet_about_the_export() -> None:
    workbook = load_workbook(io.BytesIO(exports.render_xlsx(_export())))
    assert workbook.sheetnames == ["Acknowledgements", "Export"]

    table = [[cell.value or "" for cell in row] for row in workbook["Acknowledgements"].iter_rows()]
    assert table == exports.table(_export())

    facts = {row[0].value: row[1].value for row in workbook["Export"].iter_rows()}
    assert facts["Exported by"] == "Ada Admin (ada@example.test)"
    assert facts["Exported at"] == "2026-10-06T09:30:15Z"
    assert facts["Campaign"] == "Annual acknowledgement"
    assert facts["Campaign status"] == "Active"
    assert (facts["People asked"], facts["Signed"], facts["Pending"]) == ("2", "1", "1")


def test_no_cell_in_the_excel_file_is_a_formula() -> None:
    """openpyxl stores a string that starts with ``=`` as a live formula, so the guard
    has to run before the write, not on the way out."""
    export = _export(
        campaign_title="=1+1",
        exported_by="=cmd|' /C calc'!A0",
        recipients=[
            exports.ExportRecipient(
                name="=NAME()",
                email="+1@example.test",
                kind="reviewer",
                signed=False,
                acknowledged_at=None,
                comment="@SUM(1)",
            )
        ],
    )
    workbook = load_workbook(io.BytesIO(exports.render_xlsx(export)))
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f", f"{sheet.title}!{cell.coordinate} is a formula"
    assert workbook["Acknowledgements"]["F2"].value == "'=NAME()"


def test_the_file_name_is_plain_ascii_and_dated() -> None:
    assert exports.filename(_export(), "csv") == "POL-01-acknowledgements-2026-10-06.csv"
    assert exports.filename(_export(document_code='A/B "x"'), "xlsx") == (
        "A-B-x-acknowledgements-2026-10-06.xlsx"
    )
    assert exports.filename(_export(document_code="///"), "csv").startswith("document-")
