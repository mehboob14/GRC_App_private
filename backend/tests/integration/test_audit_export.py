"""``GET /api/v1/audit-log/export`` through the real app: real tokens, permissions and RLS.

The rows an assertion cares about are seeded straight into ``audit_log`` with a chosen
``occurred_at`` and a made-up ``object_type``, then selected with the route's own
``object_type`` and date filters. Signing a workspace up writes hundreds of audit rows
of its own, so no test here counts the whole trail.
"""

from __future__ import annotations

import codecs
import csv
import io
import re
from collections.abc import AsyncIterator
from datetime import UTC, date, datetime, time, timedelta
from typing import Any, cast

import httpx
import pytest
from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy.pool import QueuePool

from tests.support.audit import full_stream
from tests.support.iam import (
    INVITEE_PASSWORD,
    Workspace,
    invite_directly,
    signup_workspace,
    tenant_session_headers,
)
from verity.core.config import Settings
from verity.core.db import dispose_engine, get_engine, session_scope
from verity.main import create_app
from verity.modules.audit import service as audit_service_module
from verity.modules.audit.export import HEADERS
from verity.modules.audit.models import AuditLog
from verity.modules.audit.schemas import iso_utc_z
from verity.modules.audit.service import ExportRequest, Membership, audit_service
from verity.modules.iam.service import iam_auth_service
from verity.shared.ids import uuid7

pytestmark = [pytest.mark.integration]

URL = "/api/v1/audit-log/export"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Seeded rows live well in the past so the cut-off never reaches them and the rows a
# signup writes (all stamped now) fall outside any date range these tests ask for.
DAY = (datetime.now(UTC) - timedelta(days=60)).date()


def _at(days: int, *, hour: int = 0, minute: int = 0, second: int = 0, micro: int = 0) -> datetime:
    moment = time(hour, minute, second, micro)
    return datetime.combine(DAY + timedelta(days=days), moment, tzinfo=UTC)


def _z(moment: datetime) -> str:
    return moment.isoformat().replace("+00:00", "Z")


def _day(days: int) -> str:
    return (DAY + timedelta(days=days)).isoformat()


@pytest.fixture(autouse=True)
async def _fresh(
    clean_iam: None, clean_tenancy: None, clean_audit_log: None
) -> AsyncIterator[None]:
    await dispose_engine()
    yield
    await dispose_engine()


@pytest.fixture
def app(settings: Settings) -> FastAPI:
    return create_app(settings)


@pytest.fixture
async def workspace() -> Workspace:
    return await signup_workspace(company="Ledger Ltd", email="founder@ledger.example")


def _client(app: FastAPI, headers: dict[str, str] | None = None) -> httpx.AsyncClient:
    return httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test", headers=headers
    )


@pytest.fixture
async def client(app: FastAPI, workspace: Workspace) -> AsyncIterator[httpx.AsyncClient]:
    headers = {"Authorization": f"Bearer {workspace.session_token}"}
    async with _client(app, headers) as http_client:
        yield http_client


async def _seed(
    workspace: Workspace,
    object_type: str,
    occurred_at: datetime,
    *,
    before: dict[str, Any] | None = None,
    after: dict[str, Any] | None = None,
) -> AuditLog:
    row = AuditLog(
        id=uuid7(),
        tenant_id=workspace.tenant_id,
        actor_type="membership",
        actor_id=workspace.membership_id,
        action="update",
        object_type=object_type,
        object_id=uuid7(),
        before=before,
        after=after,
        occurred_at=occurred_at,
    )
    async with session_scope(workspace.tenant_id) as session:
        session.add(row)
    return row


def _csv_rows(response: httpx.Response) -> list[dict[str, str]]:
    assert response.status_code == 200, response.text
    return list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))


def _sheet_rows(response: httpx.Response, sheet: str = "Audit log") -> list[tuple[Any, ...]]:
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    return [tuple(row) for row in book[sheet].iter_rows(values_only=True)]


async def _export_rows(workspace: Workspace) -> list[AuditLog]:
    stream = await full_stream(workspace.tenant_id)
    return [row for row in stream if row.object_type == "audit_export"]


# ---------------------------------------------------------------------------
# The file
# ---------------------------------------------------------------------------


async def test_csv_has_the_header_and_the_rows_oldest_first(
    client: httpx.AsyncClient, workspace: Workspace
) -> None:
    # Written newest first on purpose: the export must not depend on insert order.
    newest = await _seed(workspace, "widget", _at(2, hour=9), before={"status": "open"})
    oldest = await _seed(
        workspace,
        "widget",
        _at(0, hour=9, micro=250),
        before={"status": "draft", "owner": "ada"},
        after={"status": "open", "owner": "ada"},
    )
    middle = await _seed(workspace, "widget", _at(1, hour=9), after={"status": "closed"})

    response = await client.get(URL, params={"format": "csv", "object_type": "widget"})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert response.headers["cache-control"] == "no-store"
    assert "content-length" not in response.headers, "streamed, not built whole first"
    assert re.fullmatch(
        r'attachment; filename="audit-log-\d{4}-\d{2}-\d{2}\.csv"',
        response.headers["content-disposition"],
    )
    assert response.content.startswith(codecs.BOM_UTF8), "Excel reads UTF-8 only with a BOM"
    assert response.content.decode("utf-8-sig").splitlines()[0] == ",".join(HEADERS)

    rows = _csv_rows(response)
    assert [row["object_id"] for row in rows] == [
        str(oldest.object_id),
        str(middle.object_id),
        str(newest.object_id),
    ]
    first = rows[0]
    assert first["occurred_at"] == _z(_at(0, hour=9, micro=250))
    assert first["actor"] == "Founding Admin"
    assert first["actor_type"] == "membership"
    assert first["actor_id"] == str(workspace.membership_id)
    assert first["action"] == "update"
    assert first["object_type"] == "widget"
    assert first["object"] == "Widget", "an object with no name of its own reads as its kind"
    assert first["before"] == '{"owner":"ada","status":"draft"}'
    assert first["after"] == '{"owner":"ada","status":"open"}'
    assert rows[1]["before"] == "", "no snapshot is an empty cell, not the word null"


async def test_from_and_to_are_utc_days_and_both_are_inclusive(
    client: httpx.AsyncClient, workspace: Workspace
) -> None:
    moments = [
        _at(0, hour=23, minute=59, second=59, micro=999_999),  # last instant of day 0
        _at(1),  # first instant of day 1
        _at(1, hour=12),
        _at(2, hour=23, minute=59, second=59, micro=999_999),  # last instant of day 2
        _at(3),  # first instant of day 3
    ]
    seeded = [await _seed(workspace, "widget", moment) for moment in moments]
    ids = [str(row.object_id) for row in seeded]

    async def exported(**params: str) -> list[str]:
        response = await client.get(URL, params={"object_type": "widget", **params})
        return [row["object_id"] for row in _csv_rows(response)]

    assert await exported(**{"from": _day(1), "to": _day(2)}) == ids[1:4]
    assert await exported(**{"from": _day(2)}) == ids[3:]
    assert await exported(to=_day(1)) == ids[:3]
    assert await exported(**{"from": _day(1), "to": _day(1)}) == ids[1:3]
    assert await exported() == ids


async def test_the_object_type_filter_and_system_activity(
    client: httpx.AsyncClient, workspace: Workspace
) -> None:
    window = {"from": _day(0), "to": _day(1)}
    widget = await _seed(workspace, "widget", _at(0, hour=1))
    gadget = await _seed(workspace, "gadget", _at(0, hour=2))
    sign_in = await _seed(workspace, "session", _at(0, hour=3))

    only_gadgets = await client.get(URL, params={**window, "object_type": "gadget"})
    assert [row["object_id"] for row in _csv_rows(only_gadgets)] == [str(gadget.object_id)]

    default = await client.get(URL, params=window)
    assert [row["object_id"] for row in _csv_rows(default)] == [
        str(widget.object_id),
        str(gadget.object_id),
    ], "auth telemetry stays out unless asked for, exactly as in the list"

    with_system = await client.get(URL, params={**window, "include_system": "true"})
    assert [row["object_id"] for row in _csv_rows(with_system)] == [
        str(widget.object_id),
        str(gadget.object_id),
        str(sign_in.object_id),
    ]


async def test_a_batch_boundary_loses_and_repeats_nothing(
    client: httpx.AsyncClient, workspace: Workspace, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Five rows, three of them stamped the same microsecond: only ``id`` orders those,
    so a keyset that dropped its tie-break would lose or repeat one at a seam."""
    tie = _at(1, hour=6)
    seeded = [
        await _seed(workspace, "widget", moment)
        for moment in (_at(0, hour=6), tie, tie, tie, _at(2, hour=6))
    ]
    expected = [
        str(row.object_id) for row in sorted(seeded, key=lambda row: (row.occurred_at, row.id))
    ]

    # 1 is a seam after every row, 2 leaves a short last batch, 5 fills one batch exactly.
    for size in (1, 2, 5, 100):
        monkeypatch.setattr(audit_service_module, "EXPORT_BATCH_SIZE", size)
        csv_ids = [
            row["object_id"]
            for row in _csv_rows(await client.get(URL, params={"object_type": "widget"}))
        ]
        assert csv_ids == expected, f"csv, batch size {size}"
        sheet = _sheet_rows(
            await client.get(URL, params={"object_type": "widget", "format": "xlsx"})
        )
        assert [row[7] for row in sheet[1:]] == expected, f"xlsx, batch size {size}"


async def test_the_audit_row_is_committed_up_front_and_nothing_is_held_between_batches(
    workspace: Workspace, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A download that dies half way is still an export that happened, and a long one
    must not sit on a connection while the browser takes its time."""
    for hour in range(3):
        await _seed(workspace, "widget", _at(0, hour=hour))
    monkeypatch.setattr(audit_service_module, "EXPORT_BATCH_SIZE", 2)

    run = await audit_service.begin_export(
        tenant_id=workspace.tenant_id,
        actor=Membership(workspace.membership_id),
        request=ExportRequest("csv", object_type="widget"),
    )
    assert [row.object_id for row in await _export_rows(workspace)] == [run.export_id]

    batches = audit_service.export_batches(run)
    rows, _names = await anext(batches)
    assert len(rows) == 2, "the first of two batches"
    assert cast("QueuePool", get_engine().pool).checkedout() == 0
    await batches.aclose()
    assert [row.object_id for row in await _export_rows(workspace)] == [run.export_id]


# ---------------------------------------------------------------------------
# Who may read it, and whose rows
# ---------------------------------------------------------------------------


async def test_a_member_without_audit_read_is_refused(app: FastAPI, workspace: Workspace) -> None:
    invited = await invite_directly(
        workspace,
        email="eng@ledger.example",
        full_name="Eli Engineer",
        role_name="Engineering Lead",
    )
    await iam_auth_service.accept_invitation(
        token=invited.invite_token, full_name="Eli Engineer", password=INVITEE_PASSWORD
    )

    async with _client(app, tenant_session_headers(invited.member.membership_id)) as engineer:
        response = await engineer.get(URL)
    assert response.status_code == 403
    assert response.json()["error"]["code"] == "permission_denied"
    assert not await _export_rows(workspace), "a refused export leaves nothing to audit"

    async with _client(app) as nobody:
        assert (await nobody.get(URL)).status_code == 401


async def test_another_tenants_rows_never_reach_this_export(
    app: FastAPI, client: httpx.AsyncClient, workspace: Workspace
) -> None:
    other = await signup_workspace(company="Other Ltd", email="founder@other.example")
    mine = await _seed(workspace, "widget", _at(0, hour=1))
    theirs = await _seed(other, "secret_of_b", _at(0, hour=2), after={"price": "classified"})
    params = {
        "from": _day(0),
        "to": _day(1),
        "include_system": "true",
        "tenant_id": str(other.tenant_id),
    }

    csv_response = await client.get(URL, params=params)
    assert [row["object_id"] for row in _csv_rows(csv_response)] == [str(mine.object_id)]
    for leaked in ("secret_of_b", "classified", str(theirs.object_id), str(other.membership_id)):
        assert leaked not in csv_response.text

    xlsx_response = await client.get(URL, params={**params, "format": "xlsx"})
    cells = {str(cell) for row in _sheet_rows(xlsx_response) for cell in row}
    assert str(mine.object_id) in cells
    assert not {"secret_of_b", str(theirs.object_id)} & cells

    headers = {"Authorization": f"Bearer {other.session_token}"}
    async with _client(app, headers) as outsider:
        rows = _csv_rows(await outsider.get(URL, params=params))
    assert [row["object_id"] for row in rows] == [str(theirs.object_id)], "and the other way round"


# ---------------------------------------------------------------------------
# The export is itself an event
# ---------------------------------------------------------------------------


async def test_an_export_audits_itself_and_never_contains_its_own_row(
    client: httpx.AsyncClient, workspace: Workspace
) -> None:
    first = await client.get(URL, params={"format": "csv", "include_system": "true"})
    first_rows = _csv_rows(first)
    assert first_rows, "a fresh workspace already has a trail"
    assert not [row for row in first_rows if row["object_type"] == "audit_export"]

    (recorded,) = await _export_rows(workspace)
    assert recorded.action == "create"
    assert recorded.actor_type == "membership"
    assert recorded.actor_id == workspace.membership_id
    assert recorded.before is None
    assert recorded.after == {
        "format": "csv",
        "from": None,
        "to": None,
        "include_system": True,
        "object_type": None,
    }, "a streamed export has no row count to give"
    assert recorded.object_id.version == 7
    assert str(recorded.object_id) not in {row["object_id"] for row in first_rows}

    # The second export sees the first one's row, since that was recorded before this
    # one began, and only its own row is left out.
    second = await client.get(URL, params={"format": "xlsx", "object_type": "audit_export"})
    sheet = _sheet_rows(second)
    assert [row[7] for row in sheet[1:]] == [str(recorded.object_id)]

    latest = max(await _export_rows(workspace), key=lambda row: (row.occurred_at, row.id))
    assert latest.object_id != recorded.object_id
    assert latest.after == {
        "format": "xlsx",
        "from": None,
        "to": None,
        "include_system": False,
        "object_type": "audit_export",
        "rows": 1,
    }, "an Excel export is counted first, so its row says how many"
    assert str(latest.object_id) not in {str(cell) for row in sheet for cell in row}


# ---------------------------------------------------------------------------
# Spreadsheets run text as formulas
# ---------------------------------------------------------------------------


async def test_formulas_a_user_typed_are_neutralised_in_both_formats(
    client: httpx.AsyncClient, workspace: Workspace
) -> None:
    payload = '=HYPERLINK("http://evil.example","open me")'
    created = await client.post("/api/v1/groups", json={"name": payload})
    assert created.status_code == 201, created.text
    group_id = created.json()["id"]

    response = await client.get(URL, params={"object_type": "group"})
    (row,) = (row for row in _csv_rows(response) if row["object_id"] == group_id)
    assert row["object"] == f"'{payload}", "the label comes from the group's own name"
    assert row["after"].startswith("{"), "a snapshot is JSON text, which no spreadsheet evaluates"

    xlsx = await client.get(URL, params={"object_type": "group", "format": "xlsx"})
    events = load_workbook(io.BytesIO(xlsx.content))["Audit log"]
    (label_cell,) = (
        cells[6] for cells in events.iter_rows(min_row=2) if cells[7].value == group_id
    )
    assert (label_cell.value, label_cell.data_type) == (f"'{payload}", "s"), "never a formula"

    everything = await client.get(URL, params={"include_system": "true", "format": "xlsx"})
    dangerous = ("=", "+", "-", "@", "\t", "\r")
    for sheet_row in _sheet_rows(everything):
        for value in sheet_row:
            assert not (isinstance(value, str) and value.startswith(dangerous)), value


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


async def test_xlsx_opens_with_the_events_and_the_exports_own_details(
    client: httpx.AsyncClient, workspace: Workspace
) -> None:
    older = await _seed(workspace, "widget", _at(0, hour=9), after={"status": "open"})
    newer = await _seed(workspace, "widget", _at(1, hour=9), before={"status": "open"})

    response = await client.get(
        URL,
        params={"format": "xlsx", "object_type": "widget", "from": _day(0), "to": _day(3)},
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX
    assert response.headers["cache-control"] == "no-store"
    assert re.fullmatch(
        r'attachment; filename="audit-log-\d{4}-\d{2}-\d{2}\.xlsx"',
        response.headers["content-disposition"],
    )
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Audit log", "Export"]

    events = book["Audit log"]
    assert events.freeze_panes == "A2"
    assert events.auto_filter.ref == "A1:J3"
    rows = [tuple(row) for row in events.iter_rows(values_only=True)]
    assert rows[0] == HEADERS
    assert [row[7] for row in rows[1:]] == [str(older.object_id), str(newer.object_id)]
    assert rows[1][:2] == (_z(_at(0, hour=9)), "Founding Admin")
    assert rows[1][9] == '{"status":"open"}'
    assert rows[2][8] == '{"status":"open"}'

    (recorded,) = await _export_rows(workspace)
    details = {row[0]: row[1] for row in _sheet_rows(response, "Export")}
    assert details == {
        "Export id": str(recorded.object_id),
        "Exported by": "Founding Admin",
        "Exported at (UTC)": iso_utc_z(recorded.occurred_at),
        "Cut-off (UTC)": iso_utc_z(recorded.occurred_at),
        "From (UTC day)": _day(0),
        "To (UTC day)": _day(3),
        "Include system activity": "No",
        "Object type": "widget",
        "Rows": 2,
    }


async def test_xlsx_refuses_a_range_over_the_ceiling_and_csv_has_none(
    client: httpx.AsyncClient, workspace: Workspace, monkeypatch: pytest.MonkeyPatch
) -> None:
    for hour in range(2):
        await _seed(workspace, "widget", _at(0, hour=hour))
    await _seed(workspace, "widget", _at(1))
    monkeypatch.setattr(audit_service_module, "EXPORT_XLSX_MAX_ROWS", 2)
    params = {"object_type": "widget", "format": "xlsx"}

    refused = await client.get(URL, params=params)
    assert refused.status_code == 422
    error = refused.json()["error"]
    assert error["code"] == "invalid_input"
    assert "narrower date range" in error["message"]
    assert "2 events" in error["message"]
    assert not await _export_rows(workspace), "nothing left the building, so nothing to audit"

    narrower = await client.get(URL, params={**params, "to": _day(0)})
    assert len(_sheet_rows(narrower)) == 3, "two rows is exactly the ceiling and is allowed"

    streamed = await client.get(URL, params={"object_type": "widget", "format": "csv"})
    assert len(_csv_rows(streamed)) == 3, "CSV streams and has no ceiling"


async def test_a_range_that_runs_backwards_and_a_bad_parameter_are_refused(
    client: httpx.AsyncClient, workspace: Workspace
) -> None:
    backwards = await client.get(URL, params={"from": _day(2), "to": _day(1)})
    assert backwards.status_code == 422
    assert backwards.json()["error"]["code"] == "invalid_input"
    assert "From date is after the To date" in backwards.json()["error"]["message"]

    refused = (
        {"format": "pdf"},
        {"from": "yesterday"},
        {"to": "2026-13-45"},
        {"object_type": "x" * 101},
    )
    for params in refused:
        response = await client.get(URL, params=params)
        assert response.status_code == 422, params
        assert response.json()["error"]["code"] == "validation_error"
    assert not await _export_rows(workspace)


async def test_the_dates_span_the_end_of_time_without_overflowing(
    client: httpx.AsyncClient, workspace: Workspace
) -> None:
    only = await _seed(workspace, "widget", _at(0, hour=1))
    response = await client.get(
        URL,
        params={"object_type": "widget", "from": date.min.isoformat(), "to": date.max.isoformat()},
    )
    assert [row["object_id"] for row in _csv_rows(response)] == [str(only.object_id)]
