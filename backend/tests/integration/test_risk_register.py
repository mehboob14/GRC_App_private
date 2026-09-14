"""The risk register through the real app: real tokens, permissions and RLS.

Shallow and wide on reachability (every static route answers with its own
shape), then the flows the signed spec names: identify, score, link a control,
treat, accept with an approver who is not the requester, and reopen when the
acceptance lapses. Taxonomy, matrix, import, library and promotion follow.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import AsyncIterator
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from typing import Any, cast

import httpx
import pytest
from fastapi import FastAPI
from sqlalchemy import select, update

from tests.support.iam import (
    INVITEE_PASSWORD,
    Workspace,
    invite_directly,
    signup_workspace,
    tenant_session_headers,
)
from verity.core.config import Settings
from verity.core.db import dispose_engine, session_scope
from verity.main import create_app
from verity.modules.audit.service import Membership
from verity.modules.compliance.control_service import control_service
from verity.modules.iam.service import iam_auth_service
from verity.modules.risk.models import RiskAcceptance, RiskTemplate
from verity.modules.risk.service import risk_service
from verity.modules.vendors.models import VendorFinding
from verity.modules.vendors.service import VendorInput, vendor_service

pytestmark = [pytest.mark.integration]

API = "/api/v1/risks"


@dataclass(frozen=True, slots=True)
class Ready:
    workspace: Workspace
    approver_id: uuid.UUID
    client: httpx.AsyncClient

    @property
    def tenant_id(self) -> uuid.UUID:
        return self.workspace.tenant_id


@pytest.fixture(autouse=True)
async def _fresh(
    clean_iam: None, clean_tenancy: None, clean_audit_log: None
) -> AsyncIterator[None]:
    await dispose_engine()
    yield
    await dispose_engine()


@pytest.fixture
def app(settings: Settings) -> FastAPI:
    return create_app(settings)


@pytest.fixture
async def ready(app: FastAPI) -> AsyncIterator[Ready]:
    workspace = await signup_workspace(company="Risky Ltd", email="founder@risky.example")
    invited = await invite_directly(
        workspace, email="approver@risky.example", full_name="Avery Stone", role_name="Admin"
    )
    await iam_auth_service.accept_invitation(
        token=invited.invite_token, full_name="Avery Stone", password=INVITEE_PASSWORD
    )
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(
        transport=transport,
        base_url="http://test",
        headers={"Authorization": f"Bearer {workspace.session_token}"},
    ) as client:
        yield Ready(workspace, invited.member.membership_id, client)


async def _register(client: httpx.AsyncClient) -> dict[str, Any]:
    response = await client.get(f"{API}/registers")
    assert response.status_code == 200, response.text
    return cast(dict[str, Any], response.json()[0])


def _category(register: dict[str, Any], name: str) -> dict[str, Any]:
    return next(c for c in register["categories"] if c["name"] == name)


def _body(register: dict[str, Any], **overrides: Any) -> dict[str, Any]:  # noqa: ANN401
    operational = _category(register, "Operational")
    body: dict[str, Any] = {
        "register_id": register["id"],
        "title": "Payroll depends on one engineer",
        "category_id": operational["id"],
        "sub_category_id": operational["children"][1]["id"],
        "inherent_likelihood": 4,
        "inherent_impact": 4,
    }
    body.update(overrides)
    return body


async def _risk(client: httpx.AsyncClient, register: dict[str, Any], **kw: Any) -> dict[str, Any]:  # noqa: ANN401
    response = await client.post(API, json=_body(register, **kw))
    assert response.status_code == 201, response.text
    return cast(dict[str, Any], response.json())


# ---------------------------------------------------------------------------
# Reachability
# ---------------------------------------------------------------------------


async def test_every_static_route_answers_with_its_own_shape(ready: Ready) -> None:
    client = ready.client
    register = await _register(client)
    assert register["is_default"] is True
    assert register["likelihood_levels"] == 5
    assert [c["name"] for c in register["categories"]][:2] == ["Strategic", "Operational"]
    rid = register["id"]

    summary = (await client.get(f"{API}/summary", params={"register_id": rid})).json()
    assert summary["heatmap_inherent"] == [[0] * 5 for _ in range(5)]
    facets = (await client.get(f"{API}/facets", params={"register_id": rid})).json()
    assert set(facets) == {"statuses", "bands", "treatments", "attention"}
    page = (await client.get(API, params={"register_id": rid})).json()
    assert page == {"items": [], "total": 0}
    assert isinstance(
        (await client.get(f"{API}/library", params={"register_id": rid})).json(), list
    )
    assert (await client.get(f"{API}/refs")).json() == []
    approvers = (await client.get(f"{API}/approvers")).json()
    assert {a["name"]: a["eligible"] for a in approvers} == {
        "Avery Stone": True,
        "Founding Admin": False,
    }

    csv = await client.get(f"{API}/export", params={"register_id": rid, "format": "csv"})
    assert csv.status_code == 200
    assert csv.headers["content-type"].startswith("text/csv")
    xlsx = await client.get(f"{API}/export", params={"register_id": rid, "format": "xlsx"})
    assert xlsx.status_code == 200
    assert xlsx.content[:2] == b"PK"
    template = await client.get(f"{API}/import/template", params={"register_id": rid})
    assert template.status_code == 200
    assert template.content[:2] == b"PK"


# ---------------------------------------------------------------------------
# The lifecycle the spec names
# ---------------------------------------------------------------------------


async def test_a_risk_is_scored_controlled_treated_accepted_and_reopened_on_expiry(
    ready: Ready,
) -> None:
    client = ready.client
    register = await _register(client)
    risk = await _risk(client, register)
    assert risk["code"] == "RSK-0001"
    assert (risk["inherent_score"], risk["inherent_band"]) == (16, "critical")
    assert "no_controls" in risk["attention"]
    assert risk["next_review_on"] == (datetime.now(UTC).date() + timedelta(days=90)).isoformat()

    async with session_scope(ready.tenant_id) as session:
        control = (await control_service.list_controls(session, tenant_id=ready.tenant_id))[0]
    linked = await client.post(
        f"{API}/{risk['id']}/controls", json={"control_ids": [str(control.id)]}
    )
    assert linked.status_code == 200, linked.text
    assert linked.json()["control_count"] == 1
    assert "no_controls" not in linked.json()["attention"]

    action = await client.post(
        f"{API}/{risk['id']}/actions",
        json={"title": "Document payroll runbook", "priority": "high"},
    )
    assert action.status_code == 200, action.text
    detail = action.json()
    assert detail["status"] == "in_treatment"
    assert detail["treatment"] == "mitigate"
    assert detail["actions"][0]["title"] == "Document payroll runbook"

    rescored = await client.patch(
        f"{API}/{risk['id']}",
        json=_body(
            register,
            status="in_treatment",
            treatment="mitigate",
            residual_likelihood=2,
            residual_impact=2,
        ),
    )
    assert rescored.status_code == 200, rescored.text
    assert (rescored.json()["residual_score"], rescored.json()["residual_band"]) == (4, "low")
    assert any(e["kind"] == "scored" for e in rescored.json()["history"])

    expiry = (datetime.now(UTC).date() + timedelta(days=30)).isoformat()
    selfish = await client.post(
        f"{API}/{risk['id']}/acceptances",
        json={
            "approver_membership_id": str(ready.workspace.membership_id),
            "rationale": "Residual is low",
            "expires_on": expiry,
        },
    )
    assert selfish.status_code == 422

    requested = await client.post(
        f"{API}/{risk['id']}/acceptances",
        json={
            "approver_membership_id": str(ready.approver_id),
            "rationale": "Runbook in place and a second engineer is training.",
            "expires_on": expiry,
        },
    )
    assert requested.status_code == 200, requested.text
    acceptance = requested.json()["acceptances"][0]
    assert acceptance["status"] == "pending"
    assert "acceptance_pending" in requested.json()["attention"]

    decision_url = f"{API}/{risk['id']}/acceptances/{acceptance['id']}/decision"
    by_requester = await client.post(decision_url, json={"approve": True})
    assert by_requester.status_code == 403

    approved = await client.post(
        decision_url, json={"approve": True}, headers=tenant_session_headers(ready.approver_id)
    )
    assert approved.status_code == 200, approved.text
    assert approved.json()["status"] == "accepted"
    assert approved.json()["acceptances"][0]["status"] == "active"

    async with session_scope(ready.tenant_id) as session:
        await session.execute(
            update(RiskAcceptance)
            .where(RiskAcceptance.id == uuid.UUID(acceptance["id"]))
            .values(expires_on=date.today() - timedelta(days=1))  # noqa: DTZ011
        )
    async with session_scope(ready.tenant_id) as session:
        assert await risk_service.expire_acceptances(session, tenant_id=ready.tenant_id) == 1
    async with session_scope(ready.tenant_id) as session:
        assert await risk_service.expire_acceptances(session, tenant_id=ready.tenant_id) == 0

    reopened = (await client.get(f"{API}/{risk['id']}")).json()
    assert reopened["status"] == "open"
    assert reopened["acceptances"][0]["status"] == "expired"
    assert reopened["history"][0]["kind"] == "acceptance_expired"


async def test_closing_needs_a_reason_and_reopening_needs_a_note(ready: Ready) -> None:
    client = ready.client
    risk = await _risk(client, await _register(client))
    url = f"{API}/{risk['id']}/status"
    assert (await client.post(url, json={"status": "closed"})).status_code == 422
    closed = await client.post(url, json={"status": "closed", "note": "Service retired"})
    assert closed.json()["status"] == "closed"
    assert closed.json()["closure_justification"] == "Service retired"
    assert (await client.post(url, json={"status": "mitigated"})).status_code == 422
    assert (await client.post(url, json={"status": "open"})).status_code == 422
    reopened = await client.post(url, json={"status": "open", "note": "Service is back"})
    assert reopened.json()["status"] == "open"
    assert (await client.post(url, json={"status": "accepted"})).status_code == 422


async def test_a_subcategory_must_belong_to_its_category(ready: Ready) -> None:
    client = ready.client
    register = await _register(client)
    technology = _category(register, "Technology")
    response = await client.post(
        API, json=_body(register, sub_category_id=technology["children"][0]["id"])
    )
    assert response.status_code == 422
    partial = await client.post(API, json=_body(register, inherent_impact=None))
    assert partial.status_code == 422
    too_high = await client.post(API, json=_body(register, inherent_likelihood=6))
    assert too_high.status_code == 422


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------


async def test_the_taxonomy_renames_deletes_unused_and_archives_used(ready: Ready) -> None:
    client = ready.client
    register = await _register(client)
    financial = _category(register, "Financial")
    risk = await _risk(client, register, category_id=financial["id"], sub_category_id=None)

    tree = [
        {
            "id": c["id"],
            "name": c["name"],
            "children": [{"id": s["id"], "name": s["name"]} for s in c["children"]],
        }
        for c in register["categories"]
        if c["name"] not in ("Strategic", "Financial")
    ]
    tree[0]["name"] = "Operations"
    tree.append({"name": "People", "children": [{"name": "Attrition"}]})
    saved = await client.put(
        f"{API}/registers/{register['id']}/categories", json={"categories": tree}
    )
    assert saved.status_code == 200, saved.text
    by_name = {c["name"]: c for c in saved.json()["categories"]}
    assert "Strategic" not in by_name
    assert by_name["Financial"]["archived"] is True
    assert by_name["Operations"]["id"] == _category(register, "Operational")["id"]
    assert [s["name"] for s in by_name["People"]["children"]] == ["Attrition"]
    assert (await client.get(f"{API}/{risk['id']}")).json()["category_name"] == "Financial"

    duplicate = await client.put(
        f"{API}/registers/{register['id']}/categories",
        json={"categories": [{"name": "Same"}, {"name": "same"}]},
    )
    assert duplicate.status_code == 422


async def test_the_matrix_cannot_shrink_below_a_stored_score(ready: Ready) -> None:
    client = ready.client
    register = await _register(client)
    risk = await _risk(client, register, inherent_likelihood=5, inherent_impact=5)
    patch = {
        "name": register["name"],
        "register_type": "enterprise",
        "likelihood_levels": 3,
        "impact_levels": 3,
    }
    refused = await client.patch(f"{API}/registers/{register['id']}", json=patch)
    assert refused.status_code == 409
    await client.patch(
        f"{API}/{risk['id']}", json=_body(register, inherent_likelihood=3, inherent_impact=3)
    )
    shrunk = await client.patch(f"{API}/registers/{register['id']}", json=patch)
    assert shrunk.status_code == 200, shrunk.text
    assert [b["min_score"] for b in shrunk.json()["severity_bands"]] == [1, 2, 4, 6]
    assert len(shrunk.json()["likelihood_scale"]) == 3


async def test_a_second_register_is_configured_independently(ready: Ready) -> None:
    client = ready.client
    await _register(client)
    created = await client.post(
        f"{API}/registers",
        json={
            "name": "ISO 27001",
            "register_type": "iso_27001",
            "likelihood_levels": 4,
            "impact_levels": 4,
        },
    )
    assert created.status_code == 201, created.text
    assert created.json()["is_default"] is False
    assert created.json()["severity_bands"][-1]["min_score"] == 10
    duplicate = await client.post(f"{API}/registers", json={"name": "iso 27001"})
    assert duplicate.status_code == 409
    assert len((await client.get(f"{API}/registers")).json()) == 2


# ---------------------------------------------------------------------------
# Import, library, assist, promotion
# ---------------------------------------------------------------------------


async def test_import_previews_without_writing_and_commits_only_valid_rows(ready: Ready) -> None:
    client = ready.client
    register = await _register(client)
    csv = (
        b"Title,Category,Subcategory,Business owner,Inherent likelihood,Inherent impact,Treatment\n"
        b"Phishing,Technology,Cybersecurity,founder@risky.example,4,3,Mitigate\n"
        b"No category,,,,,,\n"
        b"Wrong sub,Technology,Liquidity,,,,\n"
        b"Ghost owner,Operational,,ghost@nowhere.example,2,2,\n"
        b"Too likely,Operational,,,9,2,\n"
    )
    preview = await client.post(
        f"{API}/import/preview",
        data={"register_id": register["id"]},
        files={"file": ("risks.csv", csv, "text/csv")},
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert (body["valid"], body["invalid"]) == (2, 3)
    rows = {r["title"]: r for r in body["rows"]}
    assert rows["Phishing"]["owner_name"] == "Founding Admin"
    assert rows["Ghost owner"]["warnings"]
    assert rows["Wrong sub"]["errors"]
    assert (await client.get(API, params={"register_id": register["id"]})).json()["total"] == 0

    valid = [
        {
            k: v
            for k, v in r.items()
            if k
            not in (
                "errors",
                "warnings",
                "category_name",
                "sub_category_name",
                "owner_name",
                "department_name",
            )
        }
        for r in body["rows"]
        if not r["errors"]
    ]
    committed = await client.post(
        f"{API}/import", json={"register_id": register["id"], "rows": valid}
    )
    assert committed.status_code == 201, committed.text
    assert committed.json() == {"created": 2, "failed": []}
    listed = (await client.get(API, params={"register_id": register["id"]})).json()
    assert listed["total"] == 2
    assert {r["origin"] for r in listed["items"]} == {"import"}


async def test_the_template_round_trips_through_the_preview(ready: Ready) -> None:
    import openpyxl  # noqa: PLC0415

    client = ready.client
    register = await _register(client)
    template = await client.get(f"{API}/import/template", params={"register_id": register["id"]})
    wb = openpyxl.load_workbook(io.BytesIO(template.content))
    assert wb.sheetnames == ["Risks", "Lists", "Guide"]
    assert wb["Lists"].sheet_state == "hidden"
    ws = wb["Risks"]
    headers = [c.value for c in ws[1]]
    row = {
        "Title *": "Template risk",
        "Category *": "Technology",
        "Subcategory": "Data",
        "Status": "In treatment",
        "Inherent likelihood": 3,
        "Inherent impact": 5,
    }
    for key, value in row.items():
        ws.cell(row=2, column=headers.index(key) + 1, value=value)
    buffer = io.BytesIO()
    wb.save(buffer)
    preview = await client.post(
        f"{API}/import/preview",
        data={"register_id": register["id"]},
        files={"file": ("filled.xlsx", buffer.getvalue(), "application/octet-stream")},
    )
    assert preview.status_code == 200, preview.text
    parsed = preview.json()["rows"][0]
    assert parsed["errors"] == []
    assert (parsed["status"], parsed["sub_category_name"], parsed["inherent_impact"]) == (
        "in_treatment",
        "Data",
        5,
    )


async def _library_loaded() -> bool:
    async with session_scope(None) as session:
        return (
            await session.execute(select(RiskTemplate).limit(1))
        ).scalar_one_or_none() is not None


async def test_library_adoption_copies_text_links_controls_and_skips_repeats(ready: Ready) -> None:
    if not await _library_loaded():
        pytest.skip("risk library not seeded; run `python -m verity.manage seed-content`")
    client = ready.client
    register = await _register(client)
    adopted = await client.post(
        f"{API}/library/adopt",
        json={"register_id": register["id"], "codes": ["RL-TEC-001", "RL-OPS-001"]},
    )
    assert adopted.status_code == 200, adopted.text
    assert adopted.json()["created"] == 2
    again = await client.post(
        f"{API}/library/adopt", json={"register_id": register["id"], "codes": ["RL-TEC-001"]}
    )
    assert again.json() == {"created": 0, "skipped": 1, "controls_linked": 0}
    library = (await client.get(f"{API}/library", params={"register_id": register["id"]})).json()
    assert {t["code"] for t in library if t["adopted"]} == {"RL-TEC-001", "RL-OPS-001"}
    risks = (await client.get(API, params={"register_id": register["id"]})).json()["items"]
    unauthorised = next(r for r in risks if r["title"].startswith("Unauthorised access"))
    assert unauthorised["category_name"] == "Technology"
    assert unauthorised["origin"] == "library"
    detail = (await client.get(f"{API}/{unauthorised['id']}")).json()
    assert detail["template_code"] == "RL-TEC-001"
    if adopted.json()["controls_linked"]:
        assert detail["controls"]


async def test_assist_drafts_from_the_library_when_no_model_is_configured(
    ready: Ready, settings: Settings
) -> None:
    if settings.ai.api_key is not None:
        pytest.skip("a model key is configured; this covers the offline path")
    if not await _library_loaded():
        pytest.skip("risk library not seeded")
    client = ready.client
    register = await _register(client)
    draft = await client.post(
        f"{API}/assist",
        json={"register_id": register["id"], "title": "Ransomware encrypts servers"},
    )
    assert draft.status_code == 200, draft.text
    body = draft.json()
    assert body["source"] == "library"
    assert body["category_id"] is not None
    assert body["inherent_likelihood"] is not None
    assert body["similar"]
    assert (await client.get(API, params={"register_id": register["id"]})).json()["total"] == 0


async def test_a_vendor_finding_is_promoted_once_and_linked_to_its_vendor(ready: Ready) -> None:
    client = ready.client
    register = await _register(client)
    actor = Membership(ready.workspace.membership_id)
    async with session_scope(ready.tenant_id) as session:
        vendor = await vendor_service.create_vendor(
            session, tenant_id=ready.tenant_id, actor=actor, data=VendorInput(name="Acme Cloud")
        )
        finding = VendorFinding(
            id=uuid.uuid4(),
            tenant_id=ready.tenant_id,
            vendor_id=vendor.id,
            title="No penetration test in 18 months",
            detail="The vendor could not provide a recent test report.",
            severity="high",
        )
        session.add(finding)
    promoted = await client.post(
        f"{API}/from-vendor-finding",
        json={"finding_id": str(finding.id), "register_id": register["id"]},
    )
    assert promoted.status_code == 201, promoted.text
    risk = promoted.json()
    assert risk["origin"] == "vendor_finding"
    assert risk["category_name"] == "Third party"
    assert [link["title"] for link in risk["links"]] == ["Acme Cloud"]
    again = await client.post(
        f"{API}/from-vendor-finding",
        json={"finding_id": str(finding.id), "register_id": register["id"]},
    )
    assert again.status_code == 409
    async with session_scope(ready.tenant_id) as session:
        view = await vendor_service.get_finding(
            session, tenant_id=ready.tenant_id, finding_id=finding.id
        )
    assert str(view.promoted_risk_id) == risk["id"]
