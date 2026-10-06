"""Acknowledgement campaigns through the real app: reminders, the overdue flag, the real
acknowledgement percentage and the completion export.

The caller is a real signed-up admin whose session is validated against the database on
every request; the teammates are real members holding only the keys each test names. The
daily reminder sweep is the worker's own function, run directly. Dates are relative to
today, and a campaign is made to have been sent or chased days ago with ``backdate``,
because the clock itself cannot be moved.
"""

from __future__ import annotations

import csv
import io
import re
import uuid
from collections.abc import AsyncIterator
from datetime import UTC, datetime, timedelta
from typing import Any

import httpx
import pytest
from fastapi import FastAPI
from openpyxl import load_workbook

from tests.support.audit import full_stream
from tests.support.documents import (
    MANAGER,
    Teammate,
    add_teammate,
    backdate,
    client,
    create_document,
    inbox,
    sign,
    start_campaign,
    today,
)
from tests.support.iam import Workspace, signup_workspace, tenant_session_headers
from verity.core.config import Settings
from verity.core.db import dispose_engine, session_scope
from verity.core.errors import PermissionDenied
from verity.main import create_app
from verity.modules.audit.service import Membership
from verity.modules.documents import exports
from verity.modules.documents.service import document_service
from verity.modules.notifications.service import notification_service
from verity.workers.tasks import _remind_pending_acknowledgements

pytestmark = pytest.mark.integration

SENT_AT = "UPDATE document_ack_campaigns SET created_at = :sent WHERE id = :id"
CHASED_AT = (
    "UPDATE document_ack_campaign_recipients SET last_reminded_at = :at, reminder_count = 1 "
    "WHERE campaign_id = :campaign AND membership_id = :member"
)
ISO_STAMP = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z$")


@pytest.fixture(autouse=True)
async def _clean_state(
    clean_iam: None, clean_tenancy: None, clean_audit_log: None
) -> AsyncIterator[None]:
    await dispose_engine()
    yield
    await dispose_engine()


@pytest.fixture
def app(settings: Settings) -> FastAPI:
    return create_app(settings)


@pytest.fixture
async def workspace() -> Workspace:
    return await signup_workspace(company="Campaigns Ltd", email="founder@campaigns.example")


@pytest.fixture
async def api(app: FastAPI, workspace: Workspace) -> AsyncIterator[httpx.AsyncClient]:
    async with client(app, tenant_session_headers(workspace.membership_id)) as admin:
        yield admin


@pytest.fixture
async def team(app: FastAPI, api: httpx.AsyncClient, workspace: Workspace) -> dict[str, Teammate]:
    """Three readers: they can open a campaign and sign it, and nothing else."""
    return {
        name: await add_teammate(
            app, api, workspace, email=f"{name.lower()}@campaigns.example", name=name
        )
        for name in ("Alice", "Bob", "Carol")
    }


async def _recipients(api: httpx.AsyncClient, campaign_id: str) -> dict[str, dict[str, Any]]:
    response = await api.get(f"/documents/campaigns/{campaign_id}")
    assert response.status_code == 200, response.text
    return {r["name"]: r for r in response.json()["recipients"]}


async def _chased(workspace: Workspace) -> list[str]:
    """The kinds of notice waiting in the email outbox, oldest first."""
    async with session_scope(workspace.tenant_id) as session:
        outbox = await notification_service.pending_emails(session, tenant_id=workspace.tenant_id)
    return [n.kind for n in outbox if n.kind == "document_ack_reminder"]


# ---------------------------------------------------------------------------
# The Remind button
# ---------------------------------------------------------------------------


async def test_a_manual_reminder_chases_only_the_unsigned_and_skips_anyone_chased_today(
    api: httpx.AsyncClient, workspace: Workspace, app: FastAPI, team: dict[str, Teammate]
) -> None:
    alice, bob, carol = team["Alice"], team["Bob"], team["Carol"]
    document = await create_document(api, workspace)
    campaign = await start_campaign(
        api, document["id"], alice, bob, carol, due=today() + timedelta(days=30)
    )
    await sign(app, alice, campaign["id"])
    remind = f"/documents/campaigns/{campaign['id']}/remind"

    first = await api.post(remind)
    assert first.status_code == 200, first.text
    assert first.json() == {"reminded": 2, "skipped": 0}

    for person, expected in ((alice, 0), (bob, 1), (carol, 1)):
        notices = await inbox(app, person.headers, "document_ack_reminder")
        assert len(notices) == expected, person.name
        for notice in notices:
            assert notice["object_type"] == "document_ack_campaign"
            assert notice["object_id"] == campaign["id"]
            assert document["code"] in notice["body"]
    people = await _recipients(api, campaign["id"])
    assert (people["Alice"]["reminder_count"], people["Alice"]["last_reminded_at"]) == (0, None)
    assert people["Bob"]["reminder_count"] == 1
    assert people["Bob"]["last_reminded_at"] is not None
    assert people["Carol"]["reminder_count"] == 1

    # Within the day nobody is chased again, however often the button is pressed.
    second = await api.post(remind)
    assert second.json() == {"reminded": 0, "skipped": 2}
    assert (await _recipients(api, campaign["id"]))["Bob"]["reminder_count"] == 1

    # More than a day on, Bob is chased and Carol, chased more recently, is not.
    await backdate(
        workspace,
        CHASED_AT,
        at=datetime.now(UTC) - timedelta(days=2),
        campaign=uuid.UUID(campaign["id"]),
        member=bob.membership_id,
    )
    third = await api.post(remind)
    assert third.json() == {"reminded": 1, "skipped": 1}
    people = await _recipients(api, campaign["id"])
    assert people["Bob"]["reminder_count"] == 2
    assert people["Carol"]["reminder_count"] == 1

    # Every reminder asked for its email copy.
    assert await _chased(workspace) == ["document_ack_reminder"] * 3


async def test_a_manual_reminder_is_audited_on_the_campaign_and_a_no_op_is_not(
    api: httpx.AsyncClient, workspace: Workspace, team: dict[str, Teammate]
) -> None:
    document = await create_document(api, workspace)
    campaign = await start_campaign(api, document["id"], team["Alice"], team["Bob"])
    remind = f"/documents/campaigns/{campaign['id']}/remind"

    assert (await api.post(remind)).json() == {"reminded": 2, "skipped": 0}
    assert (await api.post(remind)).json() == {"reminded": 0, "skipped": 2}

    reminders = [
        e
        for e in await full_stream(workspace.tenant_id)
        if e.object_type == "document_ack_campaign"
        and str(e.object_id) == campaign["id"]
        and e.action == "update"
    ]
    assert len(reminders) == 1
    assert reminders[0].after == {"reminded": 2, "skipped": 0, "trigger": "manual"}
    assert reminders[0].actor_type == "membership"
    assert reminders[0].actor_id == workspace.membership_id


async def test_only_the_sender_or_a_document_manager_may_remind(
    api: httpx.AsyncClient, workspace: Workspace, app: FastAPI, team: dict[str, Teammate]
) -> None:
    alice, bob = team["Alice"], team["Bob"]
    manager = await add_teammate(
        app, api, workspace, email="manager@campaigns.example", name="Manager", keys=MANAGER
    )
    document = await create_document(api, workspace)
    campaign = await start_campaign(api, document["id"], alice)
    remind = f"/documents/campaigns/{campaign['id']}/remind"

    # A reader who did not send it and manages nothing is refused, with a reason.
    async with client(app, bob.headers) as reader:
        refused = await reader.post(remind)
    assert refused.status_code == 403
    assert refused.json()["error"]["code"] == "permission_denied"
    assert "sent this request" in refused.json()["error"]["message"]
    assert await inbox(app, alice.headers, "document_ack_reminder") == []

    # A manager who did not send it may.
    async with client(app, manager.headers) as managing:
        allowed = await managing.post(remind)
    assert allowed.status_code == 200, allowed.text
    assert allowed.json() == {"reminded": 1, "skipped": 0}

    # The sender may with no key beyond reading: ownership is enough for the service.
    mine = await start_campaign(api, document["id"], alice, title="Sent by the admin")
    async with session_scope(workspace.tenant_id) as session:
        result = await document_service.remind_campaign(
            session,
            tenant_id=workspace.tenant_id,
            actor=Membership(workspace.membership_id),
            campaign_id=uuid.UUID(mine["id"]),
            is_manager=False,
        )
    assert (result.reminded, result.skipped) == (1, 0)
    async with session_scope(workspace.tenant_id) as session:
        with pytest.raises(PermissionDenied):
            await document_service.remind_campaign(
                session,
                tenant_id=workspace.tenant_id,
                actor=Membership(bob.membership_id),
                campaign_id=uuid.UUID(mine["id"]),
                is_manager=False,
            )


async def test_a_closed_or_missing_campaign_cannot_be_chased(
    api: httpx.AsyncClient, workspace: Workspace, app: FastAPI, team: dict[str, Teammate]
) -> None:
    document = await create_document(api, workspace)
    campaign = await start_campaign(api, document["id"], team["Alice"])
    closed = await api.post(f"/documents/campaigns/{campaign['id']}/close")
    assert closed.status_code == 200, closed.text

    refused = await api.post(f"/documents/campaigns/{campaign['id']}/remind")
    assert refused.status_code == 422
    assert "closed" in refused.json()["error"]["message"]

    missing = await api.post(f"/documents/campaigns/{uuid.uuid4()}/remind")
    assert missing.status_code == 404
    assert await inbox(app, team["Alice"].headers, "document_ack_reminder") == []


# ---------------------------------------------------------------------------
# Overdue
# ---------------------------------------------------------------------------


async def test_a_campaign_is_overdue_once_its_due_day_has_passed_with_people_still_pending(
    api: httpx.AsyncClient, workspace: Workspace, app: FastAPI, team: dict[str, Teammate]
) -> None:
    alice, bob = team["Alice"], team["Bob"]
    document = await create_document(api, workspace)
    t = today()
    late = await start_campaign(api, document["id"], alice, due=t - timedelta(days=1), title="Late")
    await start_campaign(api, document["id"], alice, due=t, title="Due today")
    await start_campaign(api, document["id"], alice, due=t + timedelta(days=5), title="Ahead")
    await start_campaign(api, document["id"], alice, title="Open ended")
    done = await start_campaign(
        api, document["id"], bob, due=t - timedelta(days=1), title="Everyone signed"
    )
    await sign(app, bob, done["id"])
    shut = await start_campaign(
        api, document["id"], alice, due=t - timedelta(days=1), title="Closed"
    )
    await api.post(f"/documents/campaigns/{shut['id']}/close")

    listed = (await api.get(f"/documents/{document['id']}/campaigns")).json()
    assert {c["title"]: c["overdue"] for c in listed} == {
        "Late": True,
        "Due today": False,
        "Ahead": False,
        "Open ended": False,
        "Everyone signed": False,
        "Closed": False,
    }
    # The campaign's own page, and the response that created it, say the same.
    assert late["overdue"] is True
    assert (await api.get(f"/documents/campaigns/{late['id']}")).json()["overdue"] is True
    assert (await api.get(f"/documents/campaigns/{shut['id']}")).json()["overdue"] is False


# ---------------------------------------------------------------------------
# The daily sweep
# ---------------------------------------------------------------------------


async def test_the_daily_sweep_chases_on_the_calendar_and_never_twice_for_one_date(
    api: httpx.AsyncClient, workspace: Workspace, app: FastAPI, team: dict[str, Teammate]
) -> None:
    alice, bob = team["Alice"], team["Bob"]
    document = await create_document(api, workspace)
    t = today()
    day = timedelta(days=1)

    lead = await start_campaign(api, document["id"], alice, due=t + 2 * day, title="Inside lead")
    early = await start_campaign(api, document["id"], alice, due=t + 10 * day, title="Too early")
    due_day = await start_campaign(api, document["id"], alice, bob, due=t, title="Due today")
    weekly = await start_campaign(api, document["id"], alice, due=t - 8 * day, title="Weekly")
    recent = await start_campaign(
        api, document["id"], alice, due=t - 8 * day, title="Chased yesterday"
    )
    undated = await start_campaign(api, document["id"], alice, title="No due date")
    closed = await start_campaign(api, document["id"], alice, due=t - 2 * day, title="Closed")
    fresh = await start_campaign(api, document["id"], alice, due=t + 2 * day, title="Sent today")
    await sign(app, bob, due_day["id"])
    await api.post(f"/documents/campaigns/{closed['id']}/close")

    # All but the last were sent ten days ago.
    ten_days_ago = datetime.now(UTC) - 10 * day
    for old in (lead, early, due_day, weekly, recent, undated, closed):
        await backdate(workspace, SENT_AT, sent=ten_days_ago, id=uuid.UUID(old["id"]))
    # The weekly one was last chased on its due date, the other yesterday.
    for campaign, when in ((weekly, t - 8 * day), (recent, t - day)):
        await backdate(
            workspace,
            CHASED_AT,
            at=datetime(when.year, when.month, when.day, 12, tzinfo=UTC),
            campaign=uuid.UUID(campaign["id"]),
            member=alice.membership_id,
        )

    assert await _remind_pending_acknowledgements() == {"reminders_sent": 3}
    # A second run, or a redelivered one, finds nobody else owed a reminder.
    assert await _remind_pending_acknowledgements() == {"reminders_sent": 0}

    chased = {n["object_id"] for n in await inbox(app, alice.headers, "document_ack_reminder")}
    assert chased == {lead["id"], due_day["id"], weekly["id"]}
    # Bob signed, so he is never chased; and nobody was chased about the rest.
    assert await inbox(app, bob.headers, "document_ack_reminder") == []
    assert not chased & {early["id"], recent["id"], undated["id"], closed["id"], fresh["id"]}
    people = await _recipients(api, lead["id"])
    assert people["Alice"]["reminder_count"] == 1
    assert people["Alice"]["last_reminded_at"] is not None

    # Each reminder is on the campaign's audit trail, written by the system.
    rows = [
        e
        for e in await full_stream(workspace.tenant_id)
        if e.object_type == "document_ack_campaign"
        and e.action == "update"
        and e.after is not None
        and e.after.get("trigger") == "schedule"
    ]
    assert {str(e.object_id) for e in rows} == {lead["id"], due_day["id"], weekly["id"]}
    assert all(e.actor_type == "system" and e.actor_id is None for e in rows)
    assert all(e.after == {"reminded": 1, "trigger": "schedule"} for e in rows if e.after)


async def test_a_reminder_by_hand_does_not_use_up_the_scheduled_one(
    api: httpx.AsyncClient, workspace: Workspace, app: FastAPI, team: dict[str, Teammate]
) -> None:
    """The calendar is dates, not a count of reminders: chasing someone early by hand
    leaves the three day reminder to come."""
    alice = team["Alice"]
    document = await create_document(api, workspace)
    campaign = await start_campaign(api, document["id"], alice, due=today() + timedelta(days=2))
    await backdate(
        workspace,
        SENT_AT,
        sent=datetime.now(UTC) - timedelta(days=12),
        id=uuid.UUID(campaign["id"]),
    )
    # Chased by hand ten days ago, long before the three day mark.
    await backdate(
        workspace,
        CHASED_AT,
        at=datetime.now(UTC) - timedelta(days=10),
        campaign=uuid.UUID(campaign["id"]),
        member=alice.membership_id,
    )

    assert await _remind_pending_acknowledgements() == {"reminders_sent": 1}

    assert (await _recipients(api, campaign["id"]))["Alice"]["reminder_count"] == 2
    assert len(await inbox(app, alice.headers, "document_ack_reminder")) == 1


# ---------------------------------------------------------------------------
# The register's percentage
# ---------------------------------------------------------------------------


async def test_the_register_percentage_is_what_the_open_campaigns_have_actually_signed(
    api: httpx.AsyncClient, workspace: Workspace, app: FastAPI, team: dict[str, Teammate]
) -> None:
    alice, bob, carol = team["Alice"], team["Bob"], team["Carol"]
    document = await create_document(api, workspace, title="With campaigns")
    bare = await create_document(api, workspace, title="No campaign")

    async def percentages() -> dict[str, float | None]:
        rows = (await api.get("/documents")).json()
        return {d["id"]: d["attestation_pct"] for d in rows}

    # Nobody was asked, so there is no percentage: not 0%.
    assert await percentages() == {document["id"]: None, bare["id"]: None}

    one = await start_campaign(api, document["id"], alice, bob, title="One")
    assert (await percentages())[document["id"]] == 0.0

    await sign(app, alice, one["id"])
    assert (await percentages())[document["id"]] == 50.0

    # A second open campaign counts too: one signature in four.
    two = await start_campaign(api, document["id"], carol, workspace.membership_id, title="Two")
    assert (await percentages())[document["id"]] == 25.0
    detail = (await api.get(f"/documents/{document['id']}")).json()
    assert (detail["acknowledged"], detail["assigned_count"]) == (1, 4)
    assert detail["attestation_pct"] == 25.0

    # A closed campaign is over: only the open one is counted.
    await api.post(f"/documents/campaigns/{one['id']}/close")
    assert (await percentages())[document["id"]] == 0.0
    await sign(app, carol, two["id"])
    assert (await percentages())[document["id"]] == 50.0

    await api.post(f"/documents/campaigns/{two['id']}/close")
    assert (await percentages())[document["id"]] is None
    assert (await percentages())[bare["id"]] is None


# ---------------------------------------------------------------------------
# The completion export
# ---------------------------------------------------------------------------


async def _signed_campaign(
    api: httpx.AsyncClient, workspace: Workspace, app: FastAPI, team: dict[str, Teammate]
) -> tuple[dict[str, Any], dict[str, Any]]:
    document = await create_document(api, workspace, title="Acceptable Use Policy")
    campaign = await start_campaign(
        api,
        document["id"],
        team["Alice"],
        approvers=(team["Bob"],),
        due=today() + timedelta(days=14),
        title="Annual acknowledgement",
    )
    await sign(app, team["Alice"], campaign["id"], comment='=HYPERLINK("http://evil.test","x")')
    return document, campaign


async def test_the_csv_export_is_a_clean_table_of_who_signed_and_who_did_not(
    api: httpx.AsyncClient, workspace: Workspace, app: FastAPI, team: dict[str, Teammate]
) -> None:
    document, campaign = await _signed_campaign(api, workspace, app, team)

    response = await api.get(
        f"/documents/campaigns/{campaign['id']}/export", params={"format": "csv"}
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/csv")
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment; filename=")
    assert f"{document['code']}-acknowledgements-" in disposition
    assert disposition.endswith('.csv"')
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[0] == list(exports.HEADERS)
    assert {len(row) for row in rows} == {len(exports.HEADERS)}
    by_name = {row[5]: row for row in rows[1:]}
    assert set(by_name) == {"Alice", "Bob"}

    alice = by_name["Alice"]
    assert alice[:5] == [
        document["code"],
        "Acceptable Use Policy",
        "0.1",
        "Annual acknowledgement",
        (today() + timedelta(days=14)).isoformat(),
    ]
    assert alice[6:9] == ["alice@campaigns.example", "reviewer", "Signed"]
    assert ISO_STAMP.match(alice[9])
    # A comment that would run as a formula arrives as text.
    assert alice[10] == '\'=HYPERLINK("http://evil.test","x")'

    bob = by_name["Bob"]
    assert bob[6:9] == ["bob@campaigns.example", "approver", "Pending"]
    assert bob[9:] == ["", ""]


async def test_the_excel_export_has_the_same_table_and_says_who_exported_it_and_when(
    api: httpx.AsyncClient, workspace: Workspace, app: FastAPI, team: dict[str, Teammate]
) -> None:
    _, campaign = await _signed_campaign(api, workspace, app, team)

    response = await api.get(
        f"/documents/campaigns/{campaign['id']}/export", params={"format": "xlsx"}
    )

    assert response.status_code == 200, response.text
    assert "spreadsheetml" in response.headers["content-type"]
    assert response.headers["content-disposition"].endswith('.xlsx"')
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Acknowledgements", "Export"]

    table = [[cell.value or "" for cell in row] for row in workbook["Acknowledgements"].iter_rows()]
    assert table[0] == list(exports.HEADERS)
    assert sorted(row[5] for row in table[1:]) == ["Alice", "Bob"]
    assert [row[8] for row in sorted(table[1:], key=lambda r: r[5])] == ["Signed", "Pending"]

    facts = {row[0].value: row[1].value for row in workbook["Export"].iter_rows()}
    assert facts["Exported by"] == f"Founding Admin ({workspace.email})"
    exported_at = datetime.strptime(facts["Exported at"], "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=UTC)
    assert abs(datetime.now(UTC) - exported_at) < timedelta(minutes=5)
    assert (facts["People asked"], facts["Signed"], facts["Pending"]) == ("2", "1", "1")
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            assert all(cell.data_type != "f" for cell in row)


async def test_the_export_needs_read_access_and_is_audited(
    api: httpx.AsyncClient, workspace: Workspace, app: FastAPI, team: dict[str, Teammate]
) -> None:
    _, campaign = await _signed_campaign(api, workspace, app, team)
    export = f"/documents/campaigns/{campaign['id']}/export"
    outsider = await add_teammate(
        app,
        api,
        workspace,
        email="outsider@campaigns.example",
        name="Outsider",
        keys=("tenant:read",),
    )

    async with client(app, outsider.headers) as nobody:
        assert (await nobody.get(export, params={"format": "csv"})).status_code == 403
    async with client(app, {}) as anonymous:
        assert (await anonymous.get(export, params={"format": "csv"})).status_code == 401
    assert (await api.get(export, params={"format": "pdf"})).status_code == 422
    assert (
        await api.get(f"/documents/campaigns/{uuid.uuid4()}/export", params={"format": "csv"})
    ).status_code == 404

    async def exports_on_trail() -> list[Any]:
        return [
            e
            for e in await full_stream(workspace.tenant_id)
            if e.object_type == "document_ack_export"
        ]

    # Nothing was taken by the refused and failed attempts.
    assert await exports_on_trail() == []

    # A reader holds the key that opens the campaign, so the export is theirs too.
    async with client(app, team["Bob"].headers) as reader:
        allowed = await reader.get(export, params={"format": "xlsx"})
    assert allowed.status_code == 200, allowed.text
    assert (await api.get(export, params={"format": "csv"})).status_code == 200

    trail = await exports_on_trail()
    assert len(trail) == 2
    assert {str(e.object_id) for e in trail} == {campaign["id"]}
    assert {e.action for e in trail} == {"create"}
    assert {e.actor_id for e in trail} == {team["Bob"].membership_id, workspace.membership_id}
    assert sorted(e.after["format"] for e in trail if e.after) == ["csv", "xlsx"]
    assert all(e.after is not None and e.after["recipients"] == 2 for e in trail)
